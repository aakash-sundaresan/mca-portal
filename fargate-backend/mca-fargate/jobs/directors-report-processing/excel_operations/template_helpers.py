"""
Excel Template Helper Functions
Functions for loading and analyzing Excel templates
"""

import boto3
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import logging

logger = logging.getLogger()
s3_client = boto3.client('s3')


def get_s3_file(bucket, key):
    """Download file from S3 and return bytes"""
    try:
        response = s3_client.get_object(Bucket=bucket, Key=key)
        file_content = response['Body'].read()
        logger.info(f"Downloaded {key}, size: {len(file_content)} bytes")
        return file_content
    except Exception as e:
        logger.error(f"Error downloading S3 object {key}: {str(e)}")
        raise


def create_template_cell_map(template_data):
    """
    Create template cell map for Bedrock analysis
    
    Args:
        template_data: Excel file as bytes
    
    Returns:
        dict: {cell_address: (value, is_protected)}
    """
    file_content = BytesIO(template_data)
    template_workbook = load_workbook(file_content, data_only=True)
    template_worksheet = template_workbook["FORM"]
    
    def cell_type_template(cell):
        """Determine if cell should be included in mapping"""
        if cell.value is None and cell.fill.start_color.rgb != 'FFE0FFFF':
            return False
        elif cell.value is None and cell.fill.start_color.rgb == 'FFE0FFFF':
            return True
        else:
            return True
    
    template_cell_map = {
        cell.coordinate: (
            str(cell.value),
            (cell.protection.locked or (cell.value is not None and cell.value != ''))
        )
        for row in template_worksheet.iter_rows()
        for cell in row
        if (
            cell_type_template(cell)
            and not template_worksheet.row_dimensions[cell.row].hidden
            and not template_worksheet.column_dimensions[get_column_letter(cell.column)].hidden
        )
    }
    
    logger.info(f"Created template cell map with {len(template_cell_map)} cells")
    
    return template_cell_map